"""End-to-end coverage of the attendee flows and the organiser console."""

import io
import tempfile
from datetime import timedelta

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.templatetags.static import static
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from .models import (
    Bookmark,
    Event,
    Feedback,
    PageView,
    Photo,
    PhotoLike,
    Poll,
    PollOption,
    TshirtOrder,
    Vote,
    category_image,
    current_festival_day,
    festival_banner,
    festival_day_choices,
)
from .views import countdown_label, view_stats


def tiny_jpeg(name="shot.jpg"):
    """A real (very small) JPEG, so ImageField validation has something to read."""
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (12, 12), (255, 42, 122)).save(buffer, format="JPEG")
    return SimpleUploadedFile(name, buffer.getvalue(), content_type="image/jpeg")


def make_event(**kwargs):
    defaults = {
        "title": "Cyber Beats Live",
        "day": "Day 1",
        "time_slot": "19:00 - 21:00",
        "location": "Main Stage",
        "category": "Music",
        "image_url": "https://example.com/stage.jpg",
    }
    defaults.update(kwargs)
    return Event.objects.create(**defaults)


class PublicAppTests(TestCase):
    def setUp(self):
        self.day1 = make_event()
        self.day2 = make_event(title="Acoustic Sunset", day="Day 2", location="Electro Dome")
        self.draft = make_event(title="Secret Set", is_published=False)
        self.photo = Photo.objects.create(
            title="Main Stage Lasers", category="Stage", image_url="https://example.com/1.jpg",
            is_approved=True, likes=5,
        )
        self.hidden_photo = Photo.objects.create(
            title="Not approved yet", category="Crowd", image_url="https://example.com/2.jpg",
        )
        self.poll = Poll.objects.create(question="Best stage?")
        self.option_a = PollOption.objects.create(poll=self.poll, text="Main", votes=3, position=1)
        self.option_b = PollOption.objects.create(poll=self.poll, text="Dome", votes=1, position=2)

    def test_home_is_the_default_tab(self):
        response = self.client.get(reverse("festival:public_app"))
        self.assertEqual(response.context["active_tab"], "home")
        for block in ["Happening soon", "Straight from the crowd", "Know before you go"]:
            self.assertContains(response, block)

    def test_home_hero_uses_latest_approved_photo(self):
        response = self.client.get(reverse("festival:public_app"))
        self.assertEqual(response.context["hero_photo"], self.photo)

    def test_home_never_shows_unapproved_photos(self):
        response = self.client.get(reverse("festival:public_app"))
        shown = [self.photo] + list(response.context["home_photos"])
        self.assertNotIn(self.hidden_photo, shown)

    @override_settings(FEST_START_DATE="2099-01-10", FEST_END_DATE="2099-01-12")
    def test_countdown_before_the_festival(self):
        self.assertTrue(countdown_label().startswith("Starts in "))

    @override_settings(FEST_START_DATE="2000-01-01", FEST_END_DATE="2000-01-03")
    def test_countdown_after_the_festival(self):
        self.assertIn("wrap", countdown_label())

    @override_settings(FEST_START_DATE="not-a-date", FEST_END_DATE="nope")
    def test_countdown_survives_a_bad_date_setting(self):
        self.assertEqual(countdown_label(), "")

    @override_settings(FEST_START_DATE="2026-09-14", FEST_END_DATE="2026-09-20")
    def test_day_tabs_are_built_from_the_festival_dates(self):
        choices = festival_day_choices()
        self.assertEqual(len(choices), 7)
        self.assertEqual(choices[0], ("Day 1", "Day 1 - Mon 14 Sep"))
        self.assertEqual(choices[-1], ("Day 7", "Day 7 - Sun 20 Sep"))

    @override_settings(FEST_START_DATE="nonsense", FEST_END_DATE="nonsense")
    def test_day_tabs_fall_back_when_dates_are_unset(self):
        self.assertEqual([v for v, _ in festival_day_choices()], ["Day 1", "Day 2", "Day 3"])

    @override_settings(FEST_START_DATE="2026-09-20", FEST_END_DATE="2026-09-14")
    def test_day_tabs_ignore_a_backwards_date_range(self):
        self.assertEqual(len(festival_day_choices()), 3)

    def test_schedule_opens_on_todays_day_during_the_festival(self):
        today = timezone.localdate()
        with override_settings(
            FEST_START_DATE=str(today - timedelta(days=2)),
            FEST_END_DATE=str(today + timedelta(days=4)),
        ):
            self.assertEqual(current_festival_day(), "Day 3")
            response = self.client.get(reverse("festival:public_app"))
            self.assertEqual(response.context["selected_day"], "Day 3")

    def test_schedule_opens_on_day_one_outside_the_festival(self):
        today = timezone.localdate()
        with override_settings(
            FEST_START_DATE=str(today + timedelta(days=30)),
            FEST_END_DATE=str(today + timedelta(days=32)),
        ):
            self.assertEqual(current_festival_day(), "Day 1")

    def test_events_tab_renders_day_one(self):
        response = self.client.get(reverse("festival:public_app"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Cyber Beats Live")
        # Day filtering belongs to the events list; the home carousel is
        # deliberately cross-day, so scope this assertion to the list itself.
        day_one = self.client.get(reverse("festival:events_partial"), {"day": "Day 1"})
        self.assertContains(day_one, "Cyber Beats Live")
        self.assertNotContains(day_one, "Acoustic Sunset")

    def test_unpublished_events_are_hidden(self):
        response = self.client.get(reverse("festival:public_app"))
        self.assertNotContains(response, "Secret Set")

    def test_event_search_and_day_filter(self):
        url = reverse("festival:events_partial")
        self.assertContains(self.client.get(url, {"day": "Day 2"}), "Acoustic Sunset")
        self.assertContains(self.client.get(url, {"day": "Day 2", "q": "dome"}), "Acoustic Sunset")
        self.assertNotContains(self.client.get(url, {"day": "Day 2", "q": "zzz"}), "Acoustic Sunset")

    def test_bookmark_toggles_on_and_off(self):
        url = reverse("festival:toggle_bookmark", args=[self.day1.id])
        first = self.client.post(url, HTTP_X_REQUESTED_WITH="fetch")
        self.assertJSONEqual(
            first.content, {"saved": True, "agenda_count": 1, "event_id": self.day1.id}
        )
        self.assertEqual(Bookmark.objects.count(), 1)

        second = self.client.post(url, HTTP_X_REQUESTED_WITH="fetch")
        self.assertJSONEqual(
            second.content, {"saved": False, "agenda_count": 0, "event_id": self.day1.id}
        )
        self.assertEqual(Bookmark.objects.count(), 0)

    def test_saved_filter_only_returns_bookmarked_events(self):
        self.client.post(
            reverse("festival:toggle_bookmark", args=[self.day1.id]), HTTP_X_REQUESTED_WITH="fetch"
        )
        response = self.client.get(reverse("festival:events_partial"), {"day": "Day 1", "saved": "1"})
        self.assertContains(response, "Cyber Beats Live")

    def test_gallery_shows_only_approved_photos(self):
        response = self.client.get(reverse("festival:gallery_partial"))
        self.assertContains(response, "Main Stage Lasers")
        self.assertNotContains(response, "Not approved yet")

    def test_gallery_category_filter(self):
        response = self.client.get(reverse("festival:gallery_partial"), {"category": "Night"})
        self.assertNotContains(response, "Main Stage Lasers")

    def test_like_is_one_per_session_and_reversible(self):
        url = reverse("festival:toggle_like", args=[self.photo.id])
        liked = self.client.post(url, HTTP_X_REQUESTED_WITH="fetch").json()
        self.assertTrue(liked["liked"])
        self.assertEqual(liked["likes"], 6)
        self.assertEqual(PhotoLike.objects.count(), 1)

        unliked = self.client.post(url, HTTP_X_REQUESTED_WITH="fetch").json()
        self.assertFalse(unliked["liked"])
        self.assertEqual(unliked["likes"], 5)
        self.assertEqual(PhotoLike.objects.count(), 0)

    def test_photo_upload_waits_for_approval(self):
        with tempfile.TemporaryDirectory() as media:
            with override_settings(MEDIA_ROOT=media):
                response = self.client.post(
                    reverse("festival:upload_photo"),
                    {"title": "My capture", "media": tiny_jpeg(), "category": "Crowd",
                     "author": "Priya"},
                    HTTP_X_REQUESTED_WITH="fetch",
                )
                self.assertEqual(response.status_code, 200)
                photo = Photo.objects.get(title="My capture")
                self.assertFalse(photo.is_approved)
                self.assertTrue(photo.image)
                self.assertNotContains(
                    self.client.get(reverse("festival:gallery_partial")), "My capture"
                )

    def test_video_upload_is_stored_as_a_clip(self):
        with tempfile.TemporaryDirectory() as media:
            with override_settings(MEDIA_ROOT=media):
                clip = SimpleUploadedFile("aarti.mp4", b"\x00\x00\x00\x18ftypmp42" + b"0" * 64,
                                          content_type="video/mp4")
                response = self.client.post(
                    reverse("festival:upload_photo"),
                    {"title": "Evening aarti", "media": clip, "category": "Night", "author": "Ravi"},
                    HTTP_X_REQUESTED_WITH="fetch",
                )
                self.assertEqual(response.status_code, 200)

                item = Photo.objects.get(title="Evening aarti")
                self.assertTrue(item.is_video)
                self.assertFalse(item.image)
                self.assertIn("aarti", item.video.name)
                self.assertEqual(item.media_url, item.video.url)

    def test_approved_video_renders_a_player_in_the_gallery(self):
        with tempfile.TemporaryDirectory() as media:
            with override_settings(MEDIA_ROOT=media):
                item = Photo.objects.create(
                    title="Dhol circle", category="Crowd", is_approved=True,
                    video=SimpleUploadedFile("dhol.mp4", b"0" * 32, content_type="video/mp4"),
                )
                response = self.client.get(reverse("festival:gallery_partial"))
                self.assertContains(response, "<video")
                self.assertContains(response, 'data-kind="video"')
                self.assertContains(response, item.video.url)

    @override_settings(FEST_MAX_VIDEO_MB=1)
    def test_oversized_clip_is_refused(self):
        # Just over the (lowered) limit, so the test stays fast.
        big = SimpleUploadedFile("long.mp4", b"0" * (1024 * 1024 + 2048), content_type="video/mp4")
        response = self.client.post(
            reverse("festival:upload_photo"),
            {"title": "Whole evening", "media": big, "category": "Night"},
            HTTP_X_REQUESTED_WITH="fetch",
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("under 1 MB", str(response.json()["errors"]))
        self.assertFalse(Photo.objects.filter(title="Whole evening").exists())

    def test_upload_without_a_name_is_credited_to_a_festival_fan(self):
        with tempfile.TemporaryDirectory() as media:
            with override_settings(MEDIA_ROOT=media):
                response = self.client.post(
                    reverse("festival:upload_photo"),
                    {"title": "Quick snap", "media": tiny_jpeg(), "category": "Crowd"},
                    HTTP_X_REQUESTED_WITH="fetch",
                )
                self.assertEqual(response.status_code, 200)
                self.assertEqual(Photo.objects.get(title="Quick snap").author, "Festival Fan")

    def test_unsupported_file_type_is_refused(self):
        response = self.client.post(
            reverse("festival:upload_photo"),
            {"title": "Programme", "category": "Stage",
             "media": SimpleUploadedFile("notes.pdf", b"%PDF-1.4", content_type="application/pdf")},
            HTTP_X_REQUESTED_WITH="fetch",
        )
        self.assertEqual(response.status_code, 400)
        self.assertFalse(Photo.objects.filter(title="Programme").exists())

    def test_photo_upload_needs_an_actual_file(self):
        # Attendees upload from their phone; a URL alone is no longer accepted.
        response = self.client.post(
            reverse("festival:upload_photo"),
            {"title": "Nothing attached", "category": "Crowd",
             "image_url": "https://example.com/x.jpg"},
            HTTP_X_REQUESTED_WITH="fetch",
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("media", response.json()["errors"])
        self.assertFalse(Photo.objects.filter(title="Nothing attached").exists())

    def test_vote_counts_once_per_session(self):
        url = reverse("festival:cast_vote", args=[self.poll.id])
        self.client.post(url, {"option": self.option_a.id}, HTTP_X_REQUESTED_WITH="fetch")
        self.option_a.refresh_from_db()
        self.assertEqual(self.option_a.votes, 4)

        # A second vote from the same session must not move the counters.
        self.client.post(url, {"option": self.option_b.id}, HTTP_X_REQUESTED_WITH="fetch")
        self.option_a.refresh_from_db()
        self.option_b.refresh_from_db()
        self.assertEqual(self.option_a.votes, 4)
        self.assertEqual(self.option_b.votes, 1)
        self.assertEqual(Vote.objects.count(), 1)

    def test_vote_response_shows_results(self):
        response = self.client.post(
            reverse("festival:cast_vote", args=[self.poll.id]),
            {"option": self.option_a.id},
            HTTP_X_REQUESTED_WITH="fetch",
        )
        self.assertContains(response, "you already voted")

    def test_closed_polls_are_not_votable(self):
        self.poll.is_active = False
        self.poll.save()
        response = self.client.post(
            reverse("festival:cast_vote", args=[self.poll.id]),
            {"option": self.option_a.id},
            HTTP_X_REQUESTED_WITH="fetch",
        )
        self.assertEqual(response.status_code, 404)

    def test_feedback_submission(self):
        response = self.client.post(
            reverse("festival:submit_feedback"),
            {"rating": 4, "category": "Food & Drinks", "comment": "More water stations please",
             "contact": "fan@example.com"},
            HTTP_X_REQUESTED_WITH="fetch",
        )
        self.assertEqual(response.status_code, 200)
        entry = Feedback.objects.get()
        self.assertEqual(entry.rating, 4)
        self.assertFalse(entry.is_resolved)

    def test_feedback_requires_a_comment(self):
        response = self.client.post(
            reverse("festival:submit_feedback"),
            {"rating": 5, "category": "General Suggestion", "comment": ""},
            HTTP_X_REQUESTED_WITH="fetch",
        )
        self.assertEqual(response.status_code, 400)
        self.assertEqual(Feedback.objects.count(), 0)


class EventArtworkTests(TestCase):
    def test_event_without_a_picture_gets_artwork_for_its_category(self):
        event = make_event(category="Sloka", image_url="")
        self.assertEqual(event.display_image, static("img/cat/sloka.png"))

    def test_an_events_own_image_url_always_wins(self):
        event = make_event(category="Dance", image_url="https://example.com/troupe.jpg")
        self.assertEqual(event.display_image, "https://example.com/troupe.jpg")

    def test_unknown_category_falls_back_to_the_festival_banner(self):
        event = make_event(category="Mystery", image_url="")
        self.assertEqual(event.display_image, festival_banner())

    def test_every_category_choice_has_artwork(self):
        for code, _ in Event.CATEGORY_CHOICES:
            with self.subTest(category=code):
                self.assertEqual(
                    category_image(code), static(f"img/cat/{code.lower()}.png")
                )


class ViewCountTests(TestCase):
    def test_opening_the_app_records_a_view_with_its_tab(self):
        self.client.get(reverse("festival:public_app"))
        self.client.get(reverse("festival:public_app"), {"tab": "gallery"})

        self.assertEqual(PageView.objects.count(), 2)
        self.assertEqual(list(PageView.objects.values_list("tab", flat=True)), ["gallery", "home"])

    def test_partials_and_actions_are_not_counted_as_views(self):
        self.client.get(reverse("festival:events_partial"))
        self.client.get(reverse("festival:gallery_partial"))
        self.client.get(reverse("festival:healthz"))
        self.assertEqual(PageView.objects.count(), 0)

    def test_repeat_visits_from_one_device_count_once_as_a_device(self):
        for _ in range(4):
            self.client.get(reverse("festival:public_app"))

        stats = view_stats()
        self.assertEqual(stats["views_total"], 4)
        self.assertEqual(stats["views_devices"], 1)
        self.assertEqual(stats["views_today"], 4)

    def test_daily_breakdown_covers_seven_days_ending_today(self):
        self.client.get(reverse("festival:public_app"))
        daily = view_stats()["views_daily"]

        self.assertEqual(len(daily), 7)
        self.assertEqual(daily[-1]["date"], timezone.localdate())
        self.assertTrue(daily[-1]["is_today"])
        self.assertEqual(daily[-1]["hits"], 1)
        self.assertEqual(daily[-1]["pct"], 100)
        self.assertEqual(sum(row["hits"] for row in daily[:-1]), 0)

    def test_tab_breakdown_percentages(self):
        for tab in ["home", "home", "polls", "events"]:
            self.client.get(reverse("festival:public_app"), {"tab": tab})

        tabs = {row["tab"]: row for row in view_stats()["views_tabs"]}
        self.assertEqual(tabs["home"]["hits"], 2)
        self.assertEqual(tabs["home"]["pct"], 50)
        self.assertEqual(tabs["polls"]["hits"], 1)

    def test_home_shows_the_public_view_counter(self):
        for _ in range(3):
            self.client.get(reverse("festival:public_app"))

        response = self.client.get(reverse("festival:public_app"))
        # The count shown is the one before this load was recorded.
        self.assertEqual(response.context["site_views"], 3)
        self.assertContains(response, "fa-eye")

    def test_console_overview_shows_the_counters(self):
        self.client.get(reverse("festival:public_app"))
        get_user_model().objects.create_user("boss", password="pw12345!", is_staff=True)
        self.client.login(username="boss", password="pw12345!")

        response = self.client.get(reverse("festival:console_overview"))
        self.assertContains(response, "App views")
        self.assertContains(response, "App views, last 7 days")
        self.assertEqual(response.context["views_total"], 1)


class TshirtOrderTests(TestCase):
    def order(self, sizes=("L", "L"), **overrides):
        """Post the form the way the page does: a count, then a size per shirt."""
        payload = {
            "flat_number": "734",
            "mobile": "98765 43210",
            "name": "Pavan",
            "quantity": len(sizes),
        }
        payload.update({f"size_{i}": size for i, size in enumerate(sizes, start=1)})
        payload.update(overrides)
        return self.client.post(
            reverse("festival:order_tshirt"), payload, HTTP_X_REQUESTED_WITH="fetch"
        )

    def test_order_is_recorded_with_the_amount_due(self):
        response = self.order(name="Pavan")
        self.assertEqual(response.status_code, 200)

        order = TshirtOrder.objects.get()  # both shirts are size L, so one row
        self.assertEqual(order.flat_number, "734")
        self.assertEqual(order.mobile, "9876543210")  # spaces stripped
        self.assertEqual(order.quantity, 2)
        self.assertEqual(order.amount, 400)  # 2 x 200
        self.assertFalse(order.is_collected)

    def test_summary_comes_back_with_the_running_total(self):
        self.order()
        response = self.order(sizes=["Kids-M"])
        self.assertContains(response, "You are on the list")
        self.assertContains(response, "3 shirts total")
        self.assertContains(response, "600")

    def test_mobile_must_be_ten_digits(self):
        response = self.order(mobile="12345")
        self.assertEqual(response.status_code, 400)
        self.assertIn("mobile", response.json()["errors"])
        self.assertEqual(TshirtOrder.objects.count(), 0)

    def test_indian_country_code_is_accepted_and_stripped(self):
        self.order(sizes=["M"], mobile="+91 98765 43210")
        self.assertEqual(TshirtOrder.objects.get().mobile, "9876543210")

    def test_quantity_is_capped(self):
        response = self.order(quantity=50)
        self.assertEqual(response.status_code, 400)
        self.assertEqual(TshirtOrder.objects.count(), 0)

    def test_each_shirt_can_be_a_different_size(self):
        self.order(sizes=["L", "Kids-S", "L", "XXL"], name="Pavan")

        rows = {row.size: row.quantity for row in TshirtOrder.objects.all()}
        self.assertEqual(rows, {"L": 2, "Kids-S": 1, "XXL": 1})  # same sizes folded together
        self.assertEqual(sum(rows.values()), 4)

    def test_repeat_orders_add_to_the_same_size_row(self):
        self.order(sizes=["M"])
        self.order(sizes=["M", "M"])

        row = TshirtOrder.objects.get(size="M")
        self.assertEqual(row.quantity, 3)
        self.assertEqual(TshirtOrder.objects.count(), 1)

    def test_flat_number_must_be_digits(self):
        response = self.order(flat_number="B-404")
        self.assertEqual(response.status_code, 400)
        self.assertIn("digits only", str(response.json()["errors"]))
        self.assertEqual(TshirtOrder.objects.count(), 0)

    def test_name_is_required(self):
        response = self.order(name="")
        self.assertEqual(response.status_code, 400)
        self.assertIn("name", str(response.json()["errors"]))
        self.assertEqual(TshirtOrder.objects.count(), 0)

    def test_a_missing_size_is_refused(self):
        response = self.client.post(
            reverse("festival:order_tshirt"),
            {"flat_number": "734", "mobile": "9876543210", "name": "Pavan",
             "quantity": 3, "size_1": "L"},
            HTTP_X_REQUESTED_WITH="fetch",
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("Pick a size for each shirt", str(response.json()["errors"]))
        self.assertEqual(TshirtOrder.objects.count(), 0)

    def test_an_invalid_size_is_refused(self):
        response = self.order(sizes=["XXXXXL"])
        self.assertEqual(response.status_code, 400)
        self.assertEqual(TshirtOrder.objects.count(), 0)

    @override_settings(FEST_TSHIRT_OPEN=False)
    def test_orders_are_refused_once_closed(self):
        response = self.order()
        self.assertEqual(response.status_code, 400)
        self.assertEqual(TshirtOrder.objects.count(), 0)
        page = self.client.get(reverse("festival:public_app"), {"tab": "tshirt"})
        self.assertContains(page, "Orders are closed")

    @override_settings(FEST_TSHIRT_PRICE=250)
    def test_price_comes_from_settings(self):
        self.order(sizes=["M", "M", "M"])
        self.assertEqual(TshirtOrder.objects.get().amount, 750)
        self.assertContains(self.client.get(reverse("festival:public_app")), "250")

    def test_a_device_only_sees_its_own_orders(self):
        self.order()
        TshirtOrder.objects.create(flat_number="1099", mobile="9000000000", size="S", quantity=5)

        response = self.client.get(reverse("festival:public_app"), {"tab": "tshirt"})
        self.assertEqual(response.context["my_shirt_count"], 2)
        self.assertNotContains(response, "1099")


class TshirtConsoleTests(TestCase):
    def setUp(self):
        get_user_model().objects.create_user("organiser", password="pw12345!", is_staff=True)
        self.client.login(username="organiser", password="pw12345!")
        TshirtOrder.objects.create(flat_number="101", mobile="9000000001", size="M", quantity=2)
        TshirtOrder.objects.create(flat_number="102", mobile="9000000002", size="M", quantity=1)
        self.big = TshirtOrder.objects.create(
            flat_number="201", mobile="9000000003", size="XL", quantity=4
        )

    def test_order_sheet_totals_and_size_breakdown(self):
        response = self.client.get(reverse("festival:console_tshirts"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["total_shirts"], 7)
        self.assertEqual(response.context["total_amount"], 1400)
        self.assertEqual(response.context["total_orders"], 3)

        sizes = {row["size"]: row["shirts"] for row in response.context["by_size"]}
        self.assertEqual(sizes, {"M": 3, "XL": 4})

    def test_marking_an_order_handed_over(self):
        self.client.post(reverse("festival:console_tshirt_action", args=[self.big.id, "toggle"]))
        self.big.refresh_from_db()
        self.assertTrue(self.big.is_collected)

        response = self.client.get(reverse("festival:console_tshirts"))
        self.assertEqual(response.context["collected_shirts"], 4)
        self.assertEqual(response.context["pending_shirts"], 3)

    def test_kids_and_adult_subtotals_appear_when_both_are_ordered(self):
        import io

        from openpyxl import load_workbook

        TshirtOrder.objects.create(flat_number="303", mobile="9000000009", size="Kids-M", quantity=2)
        book = load_workbook(io.BytesIO(self.client.get(reverse("festival:console_tshirt_export")).content))
        summary = {row[0]: row[1] for row in book["Summary"].iter_rows(values_only=True) if row[0]}

        self.assertEqual(summary["Kids subtotal"], 2)
        self.assertEqual(summary["Adult subtotal"], 7)
        self.assertEqual(summary["TOTAL SHIRTS TO PRINT"], 9)

    def test_money_sheet_totals_what_each_flat_owes(self):
        import io

        from openpyxl import load_workbook

        self.client.post(reverse("festival:console_tshirt_action", args=[self.big.id, "toggle"]))
        book = load_workbook(io.BytesIO(self.client.get(reverse("festival:console_tshirt_export")).content))
        rows = {
            str(row[0]): row
            for row in book["Money"].iter_rows(min_row=4, values_only=True)
            if row[0]
        }

        # 101: 2 shirts, none handed over, so the whole amount is outstanding.
        self.assertEqual(rows["101"][3], 2)
        self.assertEqual(rows["101"][4], 400)
        self.assertEqual(rows["101"][6], 400)

        # 201: 4 shirts handed over, nothing left to collect.
        self.assertEqual(rows["201"][5], 4)
        self.assertEqual(rows["201"][6], 0)

        total = rows["TOTAL"]
        self.assertEqual(total[3], 7)
        self.assertEqual(total[4], 1400)
        self.assertEqual(total[6], 600)  # 3 shirts still owed

    def test_export_is_an_excel_workbook_with_a_summary(self):
        import io

        from openpyxl import load_workbook

        response = self.client.get(reverse("festival:console_tshirt_export"))
        self.assertIn("spreadsheetml", response["Content-Type"])
        self.assertIn(".xlsx", response["Content-Disposition"])

        book = load_workbook(io.BytesIO(response.content))
        self.assertEqual(book.sheetnames, ["Summary", "Money", "Orders"])

        summary = {row[0]: row[1] for row in book["Summary"].iter_rows(values_only=True) if row[0]}
        # The print order, in the app's size order, with the totals beneath it.
        self.assertEqual(summary["Adult M"], 3)
        self.assertEqual(summary["Adult XL"], 4)
        self.assertEqual(summary["TOTAL SHIRTS TO PRINT"], 7)
        self.assertEqual(summary["Households ordered"], 3)
        # Money lives on its own sheet now.
        self.assertNotIn("Amount to collect", summary)
        self.assertNotIn("Price per shirt", summary)

        flats = {str(row[0]) for row in book["Orders"].iter_rows(min_row=2, values_only=True)}
        self.assertEqual(flats, {"101", "102", "201"})

    def test_order_sheet_needs_staff(self):
        self.client.logout()
        response = self.client.get(reverse("festival:console_tshirts"))
        self.assertEqual(response.status_code, 302)


class HealthCheckTests(TestCase):
    def test_healthz_reports_ok_on_a_migrated_database(self):
        response = self.client.get(reverse("festival:healthz"))
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()["ok"])

    def test_healthz_needs_no_login(self):
        self.assertEqual(self.client.get(reverse("festival:healthz")).status_code, 200)


class ConsoleAccessTests(TestCase):
    def test_console_requires_staff_login(self):
        response = self.client.get(reverse("festival:console_overview"))
        self.assertEqual(response.status_code, 302)
        self.assertIn(reverse("festival:console_login"), response.url)

    def test_signed_in_non_staff_is_rejected(self):
        get_user_model().objects.create_user("fan", password="pw12345!")
        self.client.login(username="fan", password="pw12345!")
        response = self.client.get(reverse("festival:console_overview"))
        self.assertEqual(response.status_code, 302)


class ConsoleTests(TestCase):
    def setUp(self):
        get_user_model().objects.create_user(
            "organiser", password="pw12345!", is_staff=True, is_superuser=True
        )
        self.client.login(username="organiser", password="pw12345!")
        self.event = make_event()
        self.photo = Photo.objects.create(
            title="Pending shot", category="Crowd", image_url="https://example.com/p.jpg"
        )
        self.feedback = Feedback.objects.create(rating=3, category="Food & Drinks", comment="Slow queue")

    def test_all_console_pages_load(self):
        for name in [
            "festival:console_overview",
            "festival:console_events",
            "festival:console_gallery",
            "festival:console_polls",
            "festival:console_feedback",
            "festival:console_event_create",
        ]:
            with self.subTest(page=name):
                self.assertEqual(self.client.get(reverse(name)).status_code, 200)

    def test_create_event_appears_in_public_app(self):
        response = self.client.post(
            reverse("festival:console_event_create"),
            {"title": "Drone Light Show", "day": "Day 1", "time_slot": "20:00 - 20:30",
             "location": "Main Grounds", "category": "Show", "description": "",
             "image_url": "https://example.com/drone.jpg", "is_published": "on"},
        )
        self.assertRedirects(response, reverse("festival:console_events"))
        self.assertContains(self.client.get(reverse("festival:public_app")), "Drone Light Show")

    def test_event_requires_an_image_source(self):
        response = self.client.post(
            reverse("festival:console_event_create"),
            {"title": "No art", "day": "Day 1", "time_slot": "10:00", "location": "Field",
             "category": "Other", "description": "", "image_url": "", "is_published": "on"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertFalse(Event.objects.filter(title="No art").exists())

    def test_edit_and_delete_event(self):
        self.client.post(
            reverse("festival:console_event_edit", args=[self.event.id]),
            {"title": "Renamed Set", "day": "Day 3", "time_slot": "22:00", "location": "Dome",
             "category": "Music", "description": "", "image_url": "https://example.com/a.jpg",
             "is_published": "on"},
        )
        self.event.refresh_from_db()
        self.assertEqual(self.event.title, "Renamed Set")
        self.assertEqual(self.event.day, "Day 3")

        self.client.post(reverse("festival:console_event_delete", args=[self.event.id]))
        self.assertFalse(Event.objects.filter(id=self.event.id).exists())

    def test_approving_a_photo_publishes_it(self):
        self.client.post(reverse("festival:console_photo_action", args=[self.photo.id, "approve"]))
        self.photo.refresh_from_db()
        self.assertTrue(self.photo.is_approved)
        self.assertContains(self.client.get(reverse("festival:gallery_partial")), "Pending shot")

    def test_hiding_and_deleting_a_photo(self):
        self.client.post(reverse("festival:console_photo_action", args=[self.photo.id, "approve"]))
        self.client.post(reverse("festival:console_photo_action", args=[self.photo.id, "hide"]))
        self.photo.refresh_from_db()
        self.assertFalse(self.photo.is_approved)

        self.client.post(reverse("festival:console_photo_action", args=[self.photo.id, "delete"]))
        self.assertFalse(Photo.objects.filter(id=self.photo.id).exists())

    def test_create_poll_with_options(self):
        self.client.post(
            reverse("festival:console_polls"),
            {"question": "Best food stall?", "option_1": "Tacos", "option_2": "Burgers",
             "option_3": "", "option_4": "", "is_active": "on"},
        )
        poll = Poll.objects.get(question="Best food stall?")
        self.assertEqual(poll.options.count(), 2)
        self.assertContains(self.client.get(reverse("festival:polls_partial")), "Best food stall?")

    def test_close_and_reset_poll(self):
        poll = Poll.objects.create(question="Loud enough?")
        option = PollOption.objects.create(poll=poll, text="Yes", votes=9)

        self.client.post(reverse("festival:console_poll_action", args=[poll.id, "toggle"]))
        poll.refresh_from_db()
        self.assertFalse(poll.is_active)
        self.assertNotContains(self.client.get(reverse("festival:polls_partial")), "Loud enough?")

        self.client.post(reverse("festival:console_poll_action", args=[poll.id, "reset"]))
        option.refresh_from_db()
        self.assertEqual(option.votes, 0)

    def test_feedback_resolve_toggle_and_export(self):
        self.client.post(reverse("festival:console_feedback_action", args=[self.feedback.id, "toggle"]))
        self.feedback.refresh_from_db()
        self.assertTrue(self.feedback.is_resolved)

        export = self.client.get(reverse("festival:console_feedback_export"))
        self.assertEqual(export["Content-Type"], "text/csv")
        self.assertIn("Slow queue", export.content.decode())
