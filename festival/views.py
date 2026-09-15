"""Views for the public festival app and the organiser console."""

import csv
import logging
from collections import Counter
from datetime import date, timedelta

from django.conf import settings
from django.contrib import messages
from django.contrib.auth.decorators import user_passes_test
from django.db import IntegrityError, connections, transaction
from django.db.migrations.executor import MigrationExecutor
from django.db.models import Avg, Count, F, Q, Sum, Value
from django.db.models.functions import Greatest, TruncDate
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.templatetags.static import static
from django.urls import reverse
from django.utils import timezone
from django.utils.text import slugify
from django.views.decorators.http import require_POST
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .forms import (
    EventForm,
    FeedbackForm,
    PhotoForm,
    PhotoUploadForm,
    PollForm,
    TshirtOrderForm,
)
from .models import (
    Bookmark,
    Event,
    Feedback,
    PageView,
    Photo,
    PhotoLike,
    Poll,
    PollOption,
    TshirtOrder,
    Vote,
    current_festival_day,
    festival_day_choices,
)

logger = logging.getLogger(__name__)

GALLERY_CATEGORIES = [("All", "All Photos")] + list(Photo.CATEGORY_CHOICES)

# Home screen copy. Placeholder wording - edit these two lists to match the
# real festival before you hand the link to attendees.
VIBE_CHIPS = [
    "\U0001FA94 Daily Aarti",
    "\U0001F941 Dhol Tasha",
    "\U0001F35B Mahaprasadam",
    "\U0001F3AD Cultural Nights",
    "\U0001F30A Visarjan",
]

KNOW_BEFORE = [
    {
        "icon": "fa-hands-praying",
        "colour": "fest-gold",
        "title": "Aarti timings",
        "body": "Morning and evening aarti at the Main Pandal. Check the day tab for exact times.",
    },
    {
        "icon": "fa-utensils",
        "colour": "fest-accent",
        "title": "Pooja",
        "body": (
            "Prasadam will be served in front of the mandapam, followed by pooja. "
            "Tentative time: 8:00 PM onwards."
        ),
    },
]

HELP_CARD = {
    "icon": "fa-circle-question",
    "colour": "fest-green",
    "title": "Need help?",
    "body": "For quick help, please reach out to our committee members or volunteers.",
}


# --------------------------------------------------------------------------- #
# helpers
# --------------------------------------------------------------------------- #
def session_key(request):
    """Every attendee gets a session so bookmarks, likes and votes stick."""
    if not request.session.session_key:
        request.session.create()
    return request.session.session_key


def is_ajax(request):
    return request.headers.get("x-requested-with") == "fetch"


def staff_required(view):
    return user_passes_test(
        lambda u: u.is_active and u.is_staff, login_url="festival:console_login"
    )(view)


def events_context(request):
    """Filtered event list plus the filter state the template needs."""
    day = request.GET.get("day") or current_festival_day()
    query = (request.GET.get("q") or "").strip()
    saved_only = request.GET.get("saved") == "1"
    bookmarked = set(
        Bookmark.objects.filter(session_key=session_key(request)).values_list("event_id", flat=True)
    )

    events = Event.objects.filter(is_published=True, day=day)
    if query:
        events = events.filter(
            Q(title__icontains=query)
            | Q(location__icontains=query)
            | Q(category__icontains=query)
            | Q(description__icontains=query)
        )
    if saved_only:
        events = events.filter(id__in=bookmarked)

    for event in events:
        event.is_bookmarked = event.id in bookmarked

    return {
        "events": events,
        "day_choices": festival_day_choices(),
        "selected_day": day,
        "search_query": query,
        "saved_only": saved_only,
        "agenda_count": len(bookmarked),
    }


def gallery_context(request):
    category = request.GET.get("category") or "All"
    photos = Photo.objects.filter(is_approved=True)
    if category != "All":
        photos = photos.filter(category=category)

    liked = set(
        PhotoLike.objects.filter(session_key=session_key(request)).values_list("photo_id", flat=True)
    )
    for photo in photos:
        photo.is_liked = photo.id in liked

    return {
        "photos": photos,
        "gallery_categories": GALLERY_CATEGORIES,
        "selected_category": category,
    }


def decorate_poll(poll, voted_option_id=None):
    """Attach percentages and the viewer's own choice to a poll instance."""
    options = list(poll.options.all())
    total = sum(option.votes for option in options)
    for option in options:
        option.pct = option.percentage(total)
        option.is_choice = option.id == voted_option_id
    poll.decorated_options = options
    poll.total = total
    poll.voted_option_id = voted_option_id
    poll.has_voted = voted_option_id is not None
    return poll


def polls_context(request):
    key = session_key(request)
    votes = dict(Vote.objects.filter(session_key=key).values_list("poll_id", "option_id"))
    polls = Poll.objects.filter(is_active=True).prefetch_related("options")
    return {"polls": [decorate_poll(p, votes.get(p.id)) for p in polls]}


def countdown_label():
    """Human phrasing for how far away the festival is."""
    try:
        start = date.fromisoformat(settings.FEST_START_DATE)
        end = date.fromisoformat(settings.FEST_END_DATE)
    except (ValueError, AttributeError):
        return ""

    today = timezone.localdate()
    if today < start:
        days = (start - today).days
        return "Starts tomorrow" if days == 1 else f"Starts in {days} days"
    if today <= end:
        return f"Live now - day {(today - start).days + 1}"
    return "That is a wrap. See you next year!"


def hero_image_url():
    """URL of the bundled banner, or "" when none is set or it is missing.

    Resolving it here rather than in the template means a bad path logs a
    warning and falls back, instead of 500ing on the manifest lookup.
    """
    name = getattr(settings, "FEST_HERO_IMAGE", "")
    if not name:
        return ""
    try:
        return static(name)
    except ValueError:
        logger.warning("FEST_HERO_IMAGE %r is not in the static files; ignoring it.", name)
        return ""


def home_context(request):
    """Everything the landing tab shows: hero, teasers and festival facts."""
    photos = list(Photo.objects.filter(is_approved=True)[:7])
    upcoming = list(Event.objects.filter(is_published=True)[:3])
    poll = Poll.objects.filter(is_active=True).prefetch_related("options").first()

    return {
        "site_views": PageView.objects.count(),
        "hero_image": hero_image_url(),
        "hero_photo": photos[0] if photos else None,
        "home_photos": photos[1:7],
        "home_events": upcoming,
        "home_poll": decorate_poll(poll) if poll else None,
        "fest_event_name": settings.FEST_EVENT_NAME,
        "fest_venue": settings.FEST_VENUE,
        "fest_dates": settings.FEST_DATES,
        "fest_welcome": settings.FEST_WELCOME,
        "fest_countdown": countdown_label(),
        "vibe_chips": VIBE_CHIPS,
        "know_before": KNOW_BEFORE,
        "help_card": HELP_CARD,
        "help_whatsapp": (
            f"https://wa.me/{settings.FEST_HELP_WHATSAPP}" if settings.FEST_HELP_WHATSAPP else ""
        ),
        "instagram_url": settings.FEST_INSTAGRAM_URL,
        "whatsapp_group_url": settings.FEST_WHATSAPP_GROUP_URL,
    }


def tshirt_context(request):
    """Order form state, plus what this device has already reserved."""
    mine = TshirtOrder.objects.filter(session_key=session_key(request))
    return {
        "tshirt_open": settings.FEST_TSHIRT_OPEN,
        "tshirt_price": settings.FEST_TSHIRT_PRICE,
        "tshirt_occasion": settings.FEST_TSHIRT_OCCASION,
        "tshirt_note": settings.FEST_TSHIRT_NOTE,
        "tshirt_sizes": TshirtOrder.SIZE_CHOICES,
        "my_orders": mine,
        "my_shirt_count": sum(order.quantity for order in mine),
        "my_shirt_total": sum(order.amount for order in mine),
    }


@require_POST
def order_tshirt(request):
    if not settings.FEST_TSHIRT_OPEN:
        message = "T-shirt orders are closed - the count has gone to the printer."
        if is_ajax(request):
            return JsonResponse({"ok": False, "errors": {"__all__": [message]}}, status=400)
        messages.error(request, message)
        return redirect(reverse("festival:public_app") + "?tab=tshirt")

    form = TshirtOrderForm(request.POST)
    if form.is_valid():
        rows = form.save(session_key(request))
        if is_ajax(request):
            return render(request, "public/partials/tshirt_summary.html", tshirt_context(request))
        shirts = sum(row.quantity for row in rows)
        messages.success(
            request, f"{shirts} shirt(s) reserved for flat {form.cleaned_data['flat_number']}."
        )
        return redirect(reverse("festival:public_app") + "?tab=tshirt")

    if is_ajax(request):
        return JsonResponse({"ok": False, "errors": form.errors}, status=400)
    messages.error(request, "Check the form and try again.")
    return redirect(reverse("festival:public_app") + "?tab=tshirt")


def healthz(request):
    """Deploy probe: is the database reachable, and have migrations run?

    Returns 503 with a short reason rather than a blank 500, so a failing
    deploy says what is wrong in the platform log.
    """
    try:
        connections["default"].cursor().close()
    except Exception as exc:
        logger.error("Health check failed to reach the database: %s", exc)
        return JsonResponse({"ok": False, "error": "database unreachable"}, status=503)

    executor = MigrationExecutor(connections["default"])
    missing = executor.migration_plan(executor.loader.graph.leaf_nodes())
    if missing:
        logger.error("Health check found %s unapplied migration(s).", len(missing))
        return JsonResponse(
            {"ok": False, "error": "migrations not applied", "pending": len(missing)}, status=503
        )

    return JsonResponse({"ok": True, "events": Event.objects.count()})


# --------------------------------------------------------------------------- #
# public app
# --------------------------------------------------------------------------- #
def public_app(request):
    context = {
        "active_tab": request.GET.get("tab", "home"),
        "feedback_categories": Feedback.CATEGORY_CHOICES,
        "photo_categories": Photo.CATEGORY_CHOICES,
    }
    context.update(events_context(request))
    context.update(gallery_context(request))
    context.update(polls_context(request))
    context.update(home_context(request))
    context.update(tshirt_context(request))

    PageView.objects.create(session_key=session_key(request), tab=context["active_tab"][:20])
    return render(request, "public/app.html", context)


def events_partial(request):
    return render(request, "public/partials/events_list.html", events_context(request))


def gallery_partial(request):
    return render(request, "public/partials/gallery_grid.html", gallery_context(request))


def polls_partial(request):
    return render(request, "public/partials/polls_list.html", polls_context(request))


@require_POST
def toggle_bookmark(request, pk):
    event = get_object_or_404(Event, pk=pk, is_published=True)
    key = session_key(request)
    bookmark = Bookmark.objects.filter(event=event, session_key=key).first()
    if bookmark:
        bookmark.delete()
        saved = False
    else:
        Bookmark.objects.create(event=event, session_key=key)
        saved = True

    count = Bookmark.objects.filter(session_key=key).count()
    if is_ajax(request):
        return JsonResponse({"saved": saved, "agenda_count": count, "event_id": event.id})
    return redirect(request.META.get("HTTP_REFERER", reverse("festival:public_app")))


@require_POST
def toggle_like(request, pk):
    photo = get_object_or_404(Photo, pk=pk, is_approved=True)
    key = session_key(request)
    like = PhotoLike.objects.filter(photo=photo, session_key=key).first()
    rows = Photo.objects.filter(pk=photo.pk)
    if like:
        like.delete()
        # F() keeps the counter correct when several phones tap at once.
        rows.update(likes=Greatest(F("likes") - 1, Value(0)))
        liked = False
    else:
        PhotoLike.objects.create(photo=photo, session_key=key)
        rows.update(likes=F("likes") + 1)
        liked = True
    photo.refresh_from_db(fields=["likes"])

    if is_ajax(request):
        return JsonResponse({"liked": liked, "likes": photo.likes, "photo_id": photo.id})
    return redirect(request.META.get("HTTP_REFERER", reverse("festival:public_app")))


@require_POST
def upload_photo(request):
    form = PhotoUploadForm(request.POST, request.FILES)
    if form.is_valid():
        photo = form.save(commit=False)
        photo.is_approved = settings.FEST_AUTO_APPROVE_PHOTOS
        photo.save()
        note = (
            "Photo published to the gallery!"
            if photo.is_approved
            else "Photo submitted - an organiser will approve it shortly."
        )
        if is_ajax(request):
            return JsonResponse({"ok": True, "message": note, "approved": photo.is_approved})
        messages.success(request, note)
        return redirect(reverse("festival:public_app") + "?tab=gallery")

    if is_ajax(request):
        return JsonResponse({"ok": False, "errors": form.errors}, status=400)
    messages.error(request, "Could not upload that photo. Check the form and try again.")
    return redirect(reverse("festival:public_app") + "?tab=gallery")


@require_POST
def cast_vote(request, pk):
    poll = get_object_or_404(Poll, pk=pk, is_active=True)
    option = get_object_or_404(PollOption, pk=request.POST.get("option"), poll=poll)
    key = session_key(request)

    try:
        # The savepoint keeps an outer transaction usable if the unique
        # constraint (one vote per poll per session) rejects this insert.
        with transaction.atomic():
            Vote.objects.create(poll=poll, option=option, session_key=key)
    except IntegrityError:
        # Already voted on this poll - just show them the standings again.
        existing = Vote.objects.get(poll=poll, session_key=key)
        if is_ajax(request):
            return render(
                request,
                "public/partials/poll_card.html",
                {"poll": decorate_poll(poll, existing.option_id)},
            )
        return redirect(reverse("festival:public_app") + "?tab=polls")

    PollOption.objects.filter(pk=option.pk).update(votes=F("votes") + 1)
    option.refresh_from_db(fields=["votes"])

    if is_ajax(request):
        return render(
            request, "public/partials/poll_card.html", {"poll": decorate_poll(poll, option.id)}
        )
    return redirect(reverse("festival:public_app") + "?tab=polls")


@require_POST
def submit_feedback(request):
    form = FeedbackForm(request.POST)
    if form.is_valid():
        form.save()
        if is_ajax(request):
            return JsonResponse({"ok": True})
        messages.success(request, "Thank you! Your feedback reached the organisers.")
        return redirect(reverse("festival:public_app") + "?tab=feedback")

    if is_ajax(request):
        return JsonResponse({"ok": False, "errors": form.errors}, status=400)
    messages.error(request, "Please add a comment before submitting.")
    return redirect(reverse("festival:public_app") + "?tab=feedback")


# --------------------------------------------------------------------------- #
# organiser console
# --------------------------------------------------------------------------- #
def view_stats():
    """Visit counters for the console: totals, today, and a daily breakdown."""
    today = timezone.localdate()
    week_start = today - timedelta(days=6)

    per_day = {
        row["day"]: row["hits"]
        for row in PageView.objects.filter(created_at__date__gte=week_start)
        .annotate(day=TruncDate("created_at"))
        .values("day")
        .annotate(hits=Count("id"))
    }
    daily = [
        {"date": week_start + timedelta(days=offset), "hits": per_day.get(week_start + timedelta(days=offset), 0)}
        for offset in range(7)
    ]
    busiest = max((row["hits"] for row in daily), default=0)
    for row in daily:
        row["pct"] = round(row["hits"] * 100 / busiest) if busiest else 0
        row["is_today"] = row["date"] == today

    tabs = list(
        PageView.objects.exclude(tab="").values("tab").annotate(hits=Count("id")).order_by("-hits")
    )
    tab_total = sum(row["hits"] for row in tabs) or 1
    for row in tabs:
        row["pct"] = round(row["hits"] * 100 / tab_total)

    return {
        "views_total": PageView.objects.count(),
        "views_devices": PageView.objects.values("session_key").distinct().count(),
        "views_today": PageView.objects.filter(created_at__date=today).count(),
        "views_daily": daily,
        "views_tabs": tabs,
    }


@staff_required
def console_overview(request):
    photo_stats = Photo.objects.aggregate(
        total=Count("id"),
        approved=Count("id", filter=Q(is_approved=True)),
        pending=Count("id", filter=Q(is_approved=False)),
        likes=Sum("likes"),
    )
    feedback_stats = Feedback.objects.aggregate(
        total=Count("id"),
        unresolved=Count("id", filter=Q(is_resolved=False)),
        average=Avg("rating"),
    )
    context = {
        "section": "overview",
        "total_events": Event.objects.count(),
        "published_events": Event.objects.filter(is_published=True).count(),
        "photo_stats": photo_stats,
        "feedback_stats": feedback_stats,
        "active_polls": Poll.objects.filter(is_active=True).count(),
        "total_polls": Poll.objects.count(),
        "total_votes": Vote.objects.count(),
        "total_bookmarks": Bookmark.objects.count(),
        **view_stats(),
        "recent_feedback": Feedback.objects.all()[:5],
        "pending_photos": Photo.objects.filter(is_approved=False)[:4],
        "top_events": Event.objects.annotate(saves=Count("bookmarks")).order_by("-saves")[:5],
    }
    return render(request, "console/overview.html", context)


@staff_required
def console_events(request):
    events = Event.objects.annotate(saves=Count("bookmarks"))
    day = request.GET.get("day")
    if day:
        events = events.filter(day=day)
    return render(
        request,
        "console/events.html",
        {
            "section": "events",
            "events": events,
            "day_choices": festival_day_choices(),
            "selected_day": day or "",
        },
    )


@staff_required
def console_event_form(request, pk=None):
    event = get_object_or_404(Event, pk=pk) if pk else None
    if request.method == "POST":
        form = EventForm(request.POST, request.FILES, instance=event)
        if form.is_valid():
            saved = form.save()
            messages.success(request, f"Event '{saved.title}' saved.")
            return redirect("festival:console_events")
    else:
        form = EventForm(instance=event)
    return render(
        request, "console/event_form.html", {"section": "events", "form": form, "event": event}
    )


@staff_required
@require_POST
def console_event_delete(request, pk):
    event = get_object_or_404(Event, pk=pk)
    title = event.title
    event.delete()
    messages.success(request, f"Event '{title}' deleted.")
    return redirect("festival:console_events")


@staff_required
def console_gallery(request):
    form = PhotoForm()
    if request.method == "POST":
        form = PhotoForm(request.POST, request.FILES)
        if form.is_valid():
            form.save()
            messages.success(request, "Photo added to the gallery.")
            return redirect("festival:console_gallery")

    status = request.GET.get("status", "all")
    photos = Photo.objects.annotate(like_records=Count("photo_likes"))
    if status == "pending":
        photos = photos.filter(is_approved=False)
    elif status == "approved":
        photos = photos.filter(is_approved=True)

    return render(
        request,
        "console/gallery.html",
        {
            "section": "gallery",
            "photos": photos,
            "status": status,
            "form": form,
            "pending_count": Photo.objects.filter(is_approved=False).count(),
        },
    )


@staff_required
@require_POST
def console_photo_action(request, pk, action):
    photo = get_object_or_404(Photo, pk=pk)
    if action == "approve":
        photo.is_approved = True
        photo.save(update_fields=["is_approved"])
        messages.success(request, f"'{photo.title}' is now live in the gallery.")
    elif action == "hide":
        photo.is_approved = False
        photo.save(update_fields=["is_approved"])
        messages.info(request, f"'{photo.title}' hidden from the gallery.")
    elif action == "delete":
        title = photo.title
        photo.delete()
        messages.success(request, f"'{title}' deleted.")
    return redirect(request.META.get("HTTP_REFERER", reverse("festival:console_gallery")))


@staff_required
def console_polls(request):
    form = PollForm()
    if request.method == "POST":
        form = PollForm(request.POST)
        if form.is_valid():
            poll = form.save()
            messages.success(request, f"Poll '{poll.question}' published.")
            return redirect("festival:console_polls")

    polls = [decorate_poll(p) for p in Poll.objects.prefetch_related("options")]
    return render(request, "console/polls.html", {"section": "polls", "polls": polls, "form": form})


@staff_required
@require_POST
def console_poll_action(request, pk, action):
    poll = get_object_or_404(Poll, pk=pk)
    if action == "toggle":
        poll.is_active = not poll.is_active
        poll.save(update_fields=["is_active"])
        messages.info(request, "Poll opened." if poll.is_active else "Poll closed.")
    elif action == "reset":
        poll.options.update(votes=0)
        poll.votes.all().delete()
        messages.success(request, "Poll counters reset to zero.")
    elif action == "delete":
        question = poll.question
        poll.delete()
        messages.success(request, f"Poll '{question}' deleted.")
    return redirect("festival:console_polls")


@staff_required
def console_tshirts(request):
    """The order sheet: who wants what, how many of each size, and the money."""
    orders = TshirtOrder.objects.all()
    if request.GET.get("status") == "pending":
        orders = orders.filter(is_collected=False)
    elif request.GET.get("status") == "collected":
        orders = orders.filter(is_collected=True)

    totals = TshirtOrder.objects.aggregate(shirts=Sum("quantity"), orders=Count("id"))
    shirts = totals["shirts"] or 0

    by_size = list(
        TshirtOrder.objects.values("size").annotate(shirts=Sum("quantity")).order_by("-shirts")
    )
    labels = dict(TshirtOrder.SIZE_CHOICES)
    busiest = max((row["shirts"] for row in by_size), default=0)
    for row in by_size:
        row["label"] = labels.get(row["size"], row["size"])
        row["pct"] = round(row["shirts"] * 100 / busiest) if busiest else 0

    collected = TshirtOrder.objects.filter(is_collected=True).aggregate(n=Sum("quantity"))["n"] or 0

    return render(
        request,
        "console/tshirts.html",
        {
            "section": "tshirts",
            "orders": orders,
            "status": request.GET.get("status", "all"),
            "total_shirts": shirts,
            "total_orders": totals["orders"] or 0,
            "total_amount": shirts * settings.FEST_TSHIRT_PRICE,
            "collected_shirts": collected,
            "pending_shirts": shirts - collected,
            "by_size": by_size,
            "price": settings.FEST_TSHIRT_PRICE,
            "orders_open": settings.FEST_TSHIRT_OPEN,
        },
    )


@staff_required
@require_POST
def console_tshirt_action(request, pk, action):
    order = get_object_or_404(TshirtOrder, pk=pk)
    if action == "toggle":
        order.is_collected = not order.is_collected
        order.save(update_fields=["is_collected"])
    elif action == "delete":
        order.delete()
        messages.success(request, "Order removed.")
    return redirect(request.META.get("HTTP_REFERER", reverse("festival:console_tshirts")))


HEADER_FILL = PatternFill("solid", fgColor="1F2436")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
RUPEES = '"₹"#,##0'


def _style_header(sheet, row=1):
    for cell in sheet[row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def _fit_columns(sheet, widths):
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width


@staff_required
def console_tshirt_export(request):
    """Workbook with a Summary sheet for the printer and every order beside it."""
    price = settings.FEST_TSHIRT_PRICE
    orders = list(TshirtOrder.objects.all())
    shirts = sum(order.quantity for order in orders)
    collected = sum(order.quantity for order in orders if order.is_collected)

    book = Workbook()

    # ---- Summary -------------------------------------------------------
    summary = book.active
    summary.title = "Summary"
    summary["A1"] = f"{settings.FEST_BRAND_FULL} - T-shirt orders"
    summary["A1"].font = TITLE_FONT
    summary["A2"] = f"As of {timezone.localtime().strftime('%d %b %Y, %I:%M %p')}"
    summary["A2"].font = Font(italic=True, size=9)

    summary.append([])
    for label, value, fmt in [
        ("Households ordered", len({order.flat_number for order in orders}), None),
        ("Order lines", len(orders), None),
        ("Shirts to print", shirts, None),
        ("Price per shirt", price, RUPEES),
        ("Amount to collect", shirts * price, RUPEES),
        ("Handed over", collected, None),
        ("Still pending", shirts - collected, None),
    ]:
        summary.append([label, value])
        summary.cell(row=summary.max_row, column=1).font = Font(bold=True)
        if fmt:
            summary.cell(row=summary.max_row, column=2).number_format = fmt

    summary.append([])
    summary.append(["Size", "Shirts", "Amount"])
    _style_header(summary, summary.max_row)

    labels = dict(TshirtOrder.SIZE_CHOICES)
    per_size = Counter()
    for order in orders:
        per_size[order.size] += order.quantity
    # Keep the printer's list in the app's own size order, not alphabetical.
    for value, label in TshirtOrder.SIZE_CHOICES:
        if per_size.get(value):
            summary.append([label, per_size[value], per_size[value] * price])
            summary.cell(row=summary.max_row, column=3).number_format = RUPEES

    summary.append(["Total", shirts, shirts * price])
    for column in (1, 2, 3):
        summary.cell(row=summary.max_row, column=column).font = Font(bold=True)
    summary.cell(row=summary.max_row, column=3).number_format = RUPEES
    _fit_columns(summary, {"A": 22, "B": 12, "C": 14})

    # ---- Orders --------------------------------------------------------
    sheet = book.create_sheet("Orders")
    sheet.append(["Flat", "Name", "Mobile", "Size", "Quantity", "Amount", "Handed over", "Ordered"])
    _style_header(sheet)

    for order in orders:
        sheet.append(
            [
                order.flat_number,
                order.name or "",
                order.mobile,
                labels.get(order.size, order.size),
                order.quantity,
                order.amount,
                "Yes" if order.is_collected else "No",
                timezone.localtime(order.created_at).replace(tzinfo=None),
            ]
        )
        sheet.cell(row=sheet.max_row, column=6).number_format = RUPEES
        sheet.cell(row=sheet.max_row, column=8).number_format = "dd mmm yyyy hh:mm"

    _fit_columns(sheet, {"A": 10, "B": 22, "C": 15, "D": 20, "E": 10, "F": 12, "G": 13, "H": 20})
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    stamp = timezone.localtime().strftime("%Y%m%d-%H%M")
    brand = slugify(settings.FEST_BRAND) or "festival"
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{brand}-tshirts-{stamp}.xlsx"'
    book.save(response)
    return response


@staff_required
def console_feedback(request):
    entries = Feedback.objects.all()
    category = request.GET.get("category")
    status = request.GET.get("status")
    if category:
        entries = entries.filter(category=category)
    if status == "open":
        entries = entries.filter(is_resolved=False)
    elif status == "resolved":
        entries = entries.filter(is_resolved=True)

    return render(
        request,
        "console/feedback.html",
        {
            "section": "feedback",
            "entries": entries,
            "categories": Feedback.CATEGORY_CHOICES,
            "selected_category": category or "",
            "status": status or "",
            "average": Feedback.objects.aggregate(avg=Avg("rating"))["avg"],
            "open_count": Feedback.objects.filter(is_resolved=False).count(),
        },
    )


@staff_required
@require_POST
def console_feedback_action(request, pk, action):
    entry = get_object_or_404(Feedback, pk=pk)
    if action == "toggle":
        entry.is_resolved = not entry.is_resolved
        entry.save(update_fields=["is_resolved"])
    elif action == "delete":
        entry.delete()
        messages.success(request, "Feedback entry deleted.")
    return redirect(request.META.get("HTTP_REFERER", reverse("festival:console_feedback")))


@staff_required
def console_feedback_export(request):
    response = HttpResponse(content_type="text/csv")
    stamp = timezone.localtime().strftime("%Y%m%d-%H%M")
    brand = slugify(settings.FEST_BRAND) or "festival"
    response["Content-Disposition"] = f'attachment; filename="{brand}-feedback-{stamp}.csv"'
    writer = csv.writer(response)
    writer.writerow(["Submitted", "Rating", "Category", "Comment", "Contact", "Resolved"])
    for entry in Feedback.objects.all():
        writer.writerow(
            [
                timezone.localtime(entry.created_at).strftime("%Y-%m-%d %H:%M"),
                entry.rating,
                entry.category,
                entry.comment,
                entry.contact,
                "yes" if entry.is_resolved else "no",
            ]
        )
    return response
